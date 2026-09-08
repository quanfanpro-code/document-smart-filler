"""只生成独立工作簿副本，保留模板位置、已有数据和公式。"""
from __future__ import annotations

import json
import math
import os
import re
from copy import copy
from datetime import datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from pathlib import Path
from uuid import uuid4
from zipfile import BadZipFile, ZipFile

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

SOURCE_FIELDS = ["来源文件", "来源位置", "单据/事项序号", "取数说明"]
MAX_ROWS = 1048576
MAX_COLUMNS = 16384
MAX_TEXT = 32767


def _hash(path):
    with Path(path).open("rb") as stream:
        return sha256(stream.read()).hexdigest()


def inspect_template(path, sheet=None, header_row=1, start_row=None):
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("模板仅支持普通.xlsx文件，宏文件或旧格式需另存为支持的模板")
    try:
        with ZipFile(path) as package:
            unsupported = [name for name in package.namelist() if any(
                part in name.lower() for part in ("drawings/", "activex/", "ctrlprops/", "embeddings/", "externallinks/", "vbaproject"))]
        if unsupported:
            raise ValueError("模板含绘图、控件、宏或外部对象，无法保证原样保留")
        workbook = openpyxl.load_workbook(path, data_only=False, keep_links=True)
    except (BadZipFile, KeyError) as exc:
        raise ValueError("模板文件损坏或无法读取") from exc
    try:
        if sheet is None:
            if len(workbook.sheetnames) != 1:
                raise ValueError("模板包含多个工作表，请选择业务工作表")
            sheet = workbook.sheetnames[0]
        if sheet not in workbook.sheetnames:
            raise ValueError(f"模板找不到工作表：{sheet}")
        worksheet = workbook[sheet]
        header_row = int(header_row)
        if not 1 <= header_row <= MAX_ROWS:
            raise ValueError("模板表头行超出Excel范围")
        for merged in worksheet.merged_cells.ranges:
            if merged.max_row >= header_row:
                raise ValueError(f"模板表头或数据区域含合并单元格：{merged}")
        columns = {}
        for cell in worksheet[header_row]:
            if cell.value is None or not str(cell.value).strip():
                continue
            if cell.data_type == "f":
                raise ValueError(f"模板表头不能使用公式：{cell.coordinate}")
            name = str(cell.value).strip()
            if name in columns:
                raise ValueError(f"模板表头重复：{name}")
            if name in SOURCE_FIELDS or name == "核对原因":
                raise ValueError(f"模板业务字段与系统来源列重名：{name}")
            columns[name] = cell.column
        if not columns:
            raise ValueError("模板指定行没有可用表头，请重新选择表头行")
        last_column = max(columns.values())
        if last_column + len(SOURCE_FIELDS) > MAX_COLUMNS:
            raise ValueError("模板来源列超出Excel列数容量")
        for row in worksheet.iter_rows(min_row=header_row, min_col=last_column + 1):
            if any(cell.value is not None for cell in row):
                raise ValueError("模板业务表右侧存在其他内容，无法追加来源列")
        start_row = int(start_row) if start_row is not None else max(header_row + 1, worksheet.max_row + 1)
        if not header_row < start_row <= MAX_ROWS:
            raise ValueError("模板写入起始行必须位于表头下方且不超出Excel容量")
        for column in columns.values():
            cell = worksheet.cell(start_row, column)
            if cell.value is not None:
                raise ValueError(f"模板写入起始行已有数据或公式：{cell.coordinate}，请选择空白追加区域")
        return {"path": str(path.resolve()), "sheet": sheet, "header_row": header_row,
                "start_row": start_row, "fields": list(columns), "columns": columns,
                "hash": _hash(path)}
    finally:
        workbook.close()


def _text(value):
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return "\n".join(_text(item) for item in value)
    if isinstance(value, dict):
        return "\n".join(f"{key}：{_text(item)}" for key, item in value.items())
    return str(value)


def _write(cell, value, field=""):
    if isinstance(value, (dict, list, tuple)):
        value = _text(value)
    identifiers = ("编号", "号码", "字号", "识别号", "代码", "分录序号")
    if value is not None and any(word in field for word in identifiers):
        value = str(value)
    elif isinstance(value, str) and any(word in field for word in ("金额", "税额", "价税合计", "单价")) and value.strip():
        try:
            value = Decimal(value)
        except InvalidOperation:
            pass
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        number = Decimal(str(value))
        if not number.is_finite():
            raise ValueError(f"字段“{field}”金额不是有限数值")
        digits = "".join(str(digit) for digit in number.as_tuple().digits).rstrip("0")
        if len(digits) > 15:
            raise ValueError(f"字段“{field}”超过Excel的15位有效数字精度，已停止导出以免舍入")
        if isinstance(value, Decimal):
            converted = float(value)
            if math.isfinite(converted) and Decimal(str(converted)) != value:
                raise ValueError(f"字段“{field}”超出Excel可保持的数值精度")
            value = converted
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError(f"字段“{field}”数值超出Excel可保存范围")
    if isinstance(value, str):
        if len(value) > MAX_TEXT:
            raise ValueError(f"字段“{field}”文字超过Excel单元格32767字符容量，已停止导出，请查看已保存原始回答")
        if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value):
            raise ValueError(f"字段“{field}”包含Excel无法保存的控制字符，已停止导出")
        cell.value = value
        cell.data_type = "s"
        if any(word in field for word in identifiers):
            cell.number_format = "@"
    else:
        cell.value = value


def _source_values(record):
    missing = record.get("missing") or {}
    explanations = _text(missing)
    issues = _text(record.get("issues") or [])
    return {"来源文件": record.get("source_file", ""),
            "来源位置": _text(record.get("positions") or []),
            "单据/事项序号": str(record.get("document_id", "")),
            "取数说明": "\n".join(item for item in (explanations, issues) if item)}


def _header(worksheet, columns, row):
    for name, column in columns.items():
        cell = worksheet.cell(row, column)
        _write(cell, name)
        cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305D78")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        worksheet.column_dimensions[cell.column_letter].width = 22
    worksheet.freeze_panes = f"A{row + 1}"


def _fill(worksheet, records, columns, start_row, style_row=None):
    if start_row + len(records) - 1 > MAX_ROWS:
        raise ValueError("结果超出Excel工作表行数容量，已停止导出")
    styles = {column: copy(worksheet.cell(style_row, column)._style) for column in columns.values()} if style_row else {}
    for offset, record in enumerate(records):
        values = {**record.get("fields", {}), **_source_values(record),
                  "核对原因": _text(record.get("reasons") or record.get("issues") or [])}
        for field, column in columns.items():
            cell = worksheet.cell(start_row + offset, column)
            if cell.value is not None:
                raise ValueError(f"模板写入区域已有数据或公式：{worksheet.title}!{cell.coordinate}，请改用空白追加区域")
            if column in styles and offset and not cell.has_style:
                cell._style = copy(styles[column])
            _write(cell, values.get(field), field)


def export_results(records, output_dir, batch_name, fields, template=None):
    records = list(records)
    normal, review = [], []
    for record in records:
        if record.get("status") == "normal":
            normal.append(record)
        elif record.get("status") == "review":
            review.append(record)
        else:
            raise ValueError("结果尚未完成正常/人工核对分类，不能导出")
    fields = list(fields)
    if not fields or len(set(fields)) != len(fields):
        raise ValueError("业务字段不能为空或重复")
    if len(fields) + len(SOURCE_FIELDS) + 1 > MAX_COLUMNS:
        raise ValueError("结果字段超过Excel列数容量")
    if template:
        if _hash(template["path"]) != template["hash"]:
            raise ValueError("模板内容已改变，请重新识别模板或恢复原批次模板副本")
        current = inspect_template(template["path"], template["sheet"], template["header_row"], template["start_row"])
        if current["columns"] != template["columns"] or fields != current["fields"]:
            raise ValueError("模板字段或位置与批次不一致，请重新选择")
        workbook = openpyxl.load_workbook(template["path"], data_only=False, keep_links=True)
        worksheet = workbook[current["sheet"]]
        columns = dict(current["columns"])
        last_column = max(columns.values())
        for index, name in enumerate(SOURCE_FIELDS, last_column + 1):
            columns[name] = index
            _write(worksheet.cell(current["header_row"], index), name)
        start_row = current["start_row"]
        style_row = start_row
    else:
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "提取结果"
        columns = {name: index for index, name in enumerate(fields + SOURCE_FIELDS, 1)}
        _header(worksheet, columns, 1)
        start_row, style_row = 2, None
    try:
        _fill(worksheet, normal, columns, start_row, style_row)
        review_name, suffix = "需人工核对", 2
        while review_name in workbook.sheetnames:
            review_name = f"需人工核对{suffix}"
            suffix += 1
        review_sheet = workbook.create_sheet(review_name)
        review_columns = {name: index for index, name in enumerate(fields + SOURCE_FIELDS + ["核对原因"], 1)}
        _header(review_sheet, review_columns, 1)
        _fill(review_sheet, review, review_columns, 2)
        folder = Path(output_dir)
        folder.mkdir(parents=True, exist_ok=True)
        safe_name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(batch_name)).strip(" .")[:80] or "批次"
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        unique = uuid4().hex[:10]
        stage = folder / f".待发布_{unique}.xlsx"
        target = folder / f"结果_{safe_name}_{stamp}_{unique}.xlsx"
        with stage.open("xb") as stream:
            workbook.save(stream)
            stream.flush()
            os.fsync(stream.fileno())
        check = openpyxl.load_workbook(stage, data_only=False, read_only=True)
        try:
            if check.sheetnames != workbook.sheetnames:
                raise ValueError("导出回读的工作表不一致，结果尚未发布")
            for row_number in range(start_row, start_row + len(normal)):
                if check[worksheet.title].cell(row_number, columns["单据/事项序号"]).value is None:
                    raise ValueError("导出回读发现新增行缺失，结果尚未发布")
            if check[review_name].max_row != len(review) + 1:
                raise ValueError("导出回读的核对行数不一致，结果尚未发布")
        finally:
            check.close()
        stage.rename(target)
        return target
    finally:
        workbook.close()