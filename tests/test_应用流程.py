from pathlib import Path
import json
import threading
import pytest
from docx import Document
from openpyxl import load_workbook


def test_未选择记住时密钥不落盘(tmp_path):
    from src.本机设置 import load_settings, save_settings
    assert load_settings(tmp_path)['base_url'] == ''
    save_settings(tmp_path, {'base_url':'http://localhost:8000/v1','model':'演示模型','api_key':'临时密钥','remember_key':False})
    text=(tmp_path/'.local/设置.json').read_text(encoding='utf-8-sig')
    assert '临时密钥' not in text
    assert load_settings(tmp_path)['api_key'] == ''


def test_选择记住才保留且取消记住会清除(tmp_path):
    from src.本机设置 import load_settings, save_settings
    cfg={'base_url':'http://localhost:8000/v1','model':'演示模型','api_key':'本机演示密钥','remember_key':True}
    save_settings(tmp_path,cfg)
    assert load_settings(tmp_path)['api_key']=='本机演示密钥'
    cfg['remember_key']=False
    save_settings(tmp_path,cfg)
    assert load_settings(tmp_path)['api_key']==''


class ExampleClient:
    def __init__(self):
        self.calls=0
        self.attempts=[]
    def extract(self,blocks,kind,fields=None,instructions='',cancel_event=None,on_attempt=None):
        self.calls+=1
        return {'records':[{'document_id':'单据1','fields':{'票种':'增值税普通发票','发票号码':'000001','开票日期':'2026-09-08','销售方':'示例公司','购买方':'示例客户','不含税金额':'100.00','税额':'13.00','价税合计':'113.00','币种':'人民币','金额单位':'元','销售方纳税人识别号':None,'购买方纳税人识别号':None,'商品/服务明细':None},'positions':[blocks[0]['position']],'issues':[],'missing':{k:'原文未记载' for k in ['销售方纳税人识别号','购买方纳税人识别号','商品/服务明细']}}], 'complete':True,'unprocessed':[],'raw':'示例回答','usage':{}}


def make_source(path,text='示例票据，金额113元'):
    d=Document();d.add_paragraph(text);d.save(path)
    return path


def read_counts(path):
    w=load_workbook(path)
    counts={s.title:max(s.max_row-1,0) for s in w.worksheets}
    w.close()
    return counts


def test_真实读取保存导出及重启不重复(tmp_path):
    from src.应用流程 import Pipeline
    source=make_source(tmp_path/'票据.docx')
    before=source.read_bytes()
    cfg={'base_url':'http://localhost:8000/v1','api_key':'演示','model':'演示模型'}
    c=ExampleClient()
    p=Pipeline(cfg,'发票',tmp_path/'结果',tmp_path/'批次',client=c)
    result=p.run([source])
    assert result['output'] and result['complete']
    assert read_counts(result['output'])['提取结果']==1
    assert source.read_bytes()==before
    second=ExampleClient()
    r=Pipeline(cfg,'发票',tmp_path/'结果',tmp_path/'批次',client=second).run([source],resume=True)
    assert second.calls==0
    assert read_counts(r['output'])['提取结果']==1


def test_未开始即取消不请求也不丢批次(tmp_path):
    from src.应用流程 import Pipeline
    source=make_source(tmp_path/'票据.docx')
    stop=threading.Event();stop.set()
    c=ExampleClient()
    r=Pipeline({'base_url':'http://localhost/v1','api_key':'示例','model':'示例'},'发票',tmp_path/'结果',tmp_path/'批次',client=c).run([source],cancel_event=stop)
    assert c.calls==0 and r['cancelled']


def test_内容相同重复选入只处理一次(tmp_path):
    from src.应用流程 import Pipeline
    a=make_source(tmp_path/'甲.docx');b=tmp_path/'乙.docx';b.write_bytes(a.read_bytes())
    c=ExampleClient()
    r=Pipeline({'base_url':'http://localhost/v1','api_key':'示例','model':'示例'},'发票',tmp_path/'结果',tmp_path/'批次',client=c).run([a,b])
    assert c.calls==1
    assert read_counts(r['output'])['提取结果']==1

def test_中途停止后恢复包含尚未开始文件(tmp_path):
    from src.应用流程 import Pipeline
    from src.任务存储 import Store
    paths=[make_source(tmp_path/f'{i}.docx',f'不同内容{i}') for i in range(3)]
    stop=threading.Event()
    def event(e):
        if e['message'].startswith('完成当前文件'):
            stop.set()
    Pipeline({},'发票',tmp_path/'结果',tmp_path/'批次',client=ExampleClient(),on_event=event).run(paths,cancel_event=stop)
    s=Store(tmp_path/'批次');saved=s.sources();s.close()
    assert len(saved)==3
    c=ExampleClient()
    result=Pipeline({},'发票',tmp_path/'结果',tmp_path/'批次',client=c).run([],resume=True)
    assert c.calls==2 and len(result['rows'])==3


def test_源文件可与结果位于同目录(tmp_path):
    from src.应用流程 import Pipeline
    source=make_source(tmp_path/'票据.docx')
    original=source.read_bytes()
    result=Pipeline({},'发票',tmp_path,tmp_path/'批次',client=ExampleClient()).run([source])
    assert result['normal']==1 and source.read_bytes()==original


class IncompleteClient(ExampleClient):
    def extract(self,*args,**kwargs):
        answer=super().extract(*args,**kwargs)
        answer['complete']=False
        answer['unprocessed']=['尚有一部分未完成']
        return answer


def test_恢复成功后清除旧未完成提示(tmp_path):
    from src.应用流程 import Pipeline
    source=make_source(tmp_path/'票据.docx')
    Pipeline({},'发票',tmp_path/'结果',tmp_path/'批次',client=IncompleteClient()).run([source])
    result=Pipeline({},'发票',tmp_path/'结果',tmp_path/'批次',client=ExampleClient()).run([],resume=True)
    assert result['normal']==1 and result['complete']


def test_恢复读取失败仍保留已存结果(tmp_path,monkeypatch):
    from src.应用流程 import Pipeline
    source=make_source(tmp_path/'票据.docx')
    first=Pipeline({},'发票',tmp_path/'结果',tmp_path/'批次',client=IncompleteClient()).run([source])
    def fail(*args):
        raise ValueError('暂时无法读取')
    monkeypatch.setattr('src.应用流程.read_document',fail)
    second=Pipeline({},'发票',tmp_path/'结果',tmp_path/'批次',client=ExampleClient()).run([],resume=True)
    assert second['rows'][0]['fields']==first['rows'][0]['fields']
    from src.任务存储 import Store
    s=Store(tmp_path/'批次');assert s.get_checkpoint(s.sources()[0]['source_id'])['chunks'];s.close()


def test_源文件改变不能混入原批次(tmp_path):
    from src.应用流程 import Pipeline
    source=make_source(tmp_path/'票据.docx')
    Pipeline({},'发票',tmp_path/'结果',tmp_path/'批次',client=ExampleClient()).run([source])
    make_source(source,'原文发生改变')
    with pytest.raises(ValueError,match='原文件.*改变'):
        Pipeline({},'发票',tmp_path/'结果',tmp_path/'批次',client=ExampleClient()).run([source],resume=True)


def test_批次保存模板副本且恢复不受原模板改变影响(tmp_path):
    from src.应用流程 import Pipeline
    from src.模板输出 import inspect_template
    from src.任务存储 import Store
    from openpyxl import Workbook
    source=make_source(tmp_path/'票据.docx')
    template=tmp_path/'原模板.xlsx'
    w=Workbook();w.active.append(['发票号码']);w.save(template);w.close()
    info=inspect_template(template)
    first=Pipeline({},'发票',tmp_path/'结果',tmp_path/'批次',template=info,client=ExampleClient()).run([source])
    s=Store(tmp_path/'批次');saved=s.get_config()['template'];s.close()
    assert Path(saved['path'])!=template and Path(saved['path']).is_file()
    w=Workbook();w.active.append(['销售方']);w.save(template);w.close()
    second=Pipeline({},'发票',tmp_path/'结果',tmp_path/'批次',template=saved,client=ExampleClient()).run([],resume=True)
    assert second['normal']==1


def test_不同文件票号相同需提示核对且恢复保留(tmp_path):
    from src.应用流程 import Pipeline
    sources=[make_source(tmp_path/f'{i}.docx',f'不同扫描版本{i}') for i in range(2)]
    first=Pipeline({},'发票',tmp_path/'结果',tmp_path/'批次',client=ExampleClient()).run(sources)
    assert first['review']==2 and first['normal']==0
    second=Pipeline({},'发票',tmp_path/'结果',tmp_path/'批次',client=ExampleClient()).run([],resume=True)
    assert second['review']==2


def test_十页票据停止恢复结果各保留一次(tmp_path):
    import copy,re
    import pymupdf
    from src.应用流程 import Pipeline
    from src.模型调用 import Cancelled
    source=tmp_path/'十张票.pdf'
    document=pymupdf.open()
    for i in range(10):
        page=document.new_page();page.insert_text((60,60),f'Example invoice {i+1}')
    document.save(source);document.close()
    original=source.read_bytes()
    stop=threading.Event()
    class PagesClient(ExampleClient):
        def __init__(self,interrupt=False):
            super().__init__();self.interrupt=interrupt
        def extract(self,blocks,kind,fields=None,instructions='',cancel_event=None,on_attempt=None):
            if self.interrupt and blocks[0]['position']!='第1页':
                assert cancel_event.wait(5)
                raise Cancelled('测试已停止')
            base=super().extract(blocks,kind,fields,instructions,cancel_event,on_attempt)
            rows=[]
            for block in blocks:
                row=copy.deepcopy(base['records'][0])
                number=re.search(r'第(\d+)页',block['position']).group(1)
                row['document_id']='票据'+number;row['fields']['发票号码']=number.zfill(8)
                row['positions']=[block['position']];rows.append(row)
            base['records']=rows;return base
    def event(e):
        if '已保存' in e['message']:
            stop.set()
    first=Pipeline({},'发票',tmp_path/'结果',tmp_path/'批次',client=PagesClient(True),on_event=event).run([source],cancel_event=stop)
    assert first['cancelled'] and 0<len(first['rows'])<10
    second=Pipeline({},'发票',tmp_path/'结果',tmp_path/'批次',client=PagesClient()).run([],resume=True)
    assert second['normal']==10 and second['review']==0
    assert len({r['fields']['发票号码'] for r in second['rows']})==10
    assert read_counts(second['output'])['提取结果']==10 and source.read_bytes()==original
