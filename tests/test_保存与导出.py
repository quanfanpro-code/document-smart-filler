"""验证中断恢复、来源替换和真实工作簿保护，测试资料全部在临时目录生成。"""
from decimal import Decimal
from hashlib import sha256
from importlib import import_module
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Font


def storage():
    return import_module("src.任务存储").Store


def output():
    return import_module("src.模板输出")


def record(number="0001", status="normal", amount=Decimal("113.00")):
    return {
        "document_id": number,
        "fields": {"发票号码": number, "价税合计": amount, "摘要": "正常内容"},
        "positions": ["第1页"], "status": status,
        "reasons": ["金额关系不符"] if status == "review" else [],
        "issues": [], "missing": {"购买方": "原文未记载"},
        "source_file": "测试发票.pdf",
    }


def test_来源替换和重开不增加结果(tmp_path):
    Store = storage()
    batch = tmp_path / "批次"
    store = Store(batch)
    store.configure({"kind": "发票", "settings": {"model": "测试模型", "api_key": "第一次密钥"}})
    store.replace_source("source-a", "原件甲.pdf", [record(), record("0002", "review")], complete=False)
    store.replace_source("source-b", "原件乙.pdf", [record("0003")])
    store.replace_source("source-a", "原件甲.pdf", [record(), record("0002")])
    store.set_export(tmp_path / "结果.xlsx")
    store.close()
    reopened = Store(batch)
    reopened.configure({"kind": "发票", "settings": {"model": "测试模型", "api_key": "更新后的密钥"}})
    assert [r["document_id"] for r in reopened.records()] == ["0001", "0002", "0003"]
    assert [r["source_file"] for r in reopened.records()] == ["原件甲.pdf", "原件甲.pdf", "原件乙.pdf"]
    assert reopened.source_done("source-a") is True
    assert reopened.source_done("不存在") is False
    assert reopened.sources() == [
        {"source_id": "source-a", "path": "原件甲.pdf", "complete": True},
        {"source_id": "source-b", "path": "原件乙.pdf", "complete": True},
    ]
    assert reopened.get_config() == {"kind": "发票", "settings": {"model": "测试模型"}}
    assert Path(reopened.get_export()) == tmp_path / "结果.xlsx"
    reopened.close()
    for path in batch.rglob("*"):
        if path.is_file():
            assert "第一次密钥".encode() not in path.read_bytes()
            assert "更新后的密钥".encode() not in path.read_bytes()


@pytest.mark.parametrize("changed", [{"model": "另一个模型", "template": "模板甲"}, {"model": "测试模型", "template": "模板乙"}])
def test_配置改变不能复用旧批次(tmp_path, changed):
    store = storage()(tmp_path)
    store.configure({"model": "测试模型", "template": "模板甲", "api_key": "不保存"})
    with pytest.raises(ValueError, match="配置|批次"):
        store.configure(changed)
    assert store.get_config()["template"] == "模板甲"
    store.close()


def test_替换失败保留已确认结果(tmp_path):
    store = storage()(tmp_path)
    store.replace_source("source-a", "原件.pdf", [record()])
    invalid = record("0002")
    invalid["fields"]["非法对象"] = object()
    with pytest.raises((TypeError, ValueError)):
        store.replace_source("source-a", "原件.pdf", [invalid], complete=False)
    assert store.source_done("source-a") is True
    assert [r["document_id"] for r in store.records()] == ["0001"]
    store.close()


def test_每次回答分别存档且不覆盖(tmp_path):
    store = storage()(tmp_path)
    first = Path(store.save_attempt("../同名来源", "第一轮原话", {"total_tokens": 12}))
    before = first.read_bytes()
    second = Path(store.save_attempt("../同名来源", "第二轮原话", {"total_tokens": 15}))
    assert first != second and first.read_bytes() == before
    assert "第一轮原话" in first.read_text(encoding="utf-8-sig")
    assert "第二轮原话" in second.read_text(encoding="utf-8-sig")
    assert first.resolve().is_relative_to(tmp_path.resolve())
    store.close()


def test_十张票反复导出不追加并保护已改结果(tmp_path):
    module = output()
    records = [record(f"{index:04}") for index in range(1, 11)]
    first = module.export_results(records, tmp_path, "十张票", ["发票号码", "价税合计", "摘要"])
    workbook = openpyxl.load_workbook(first)
    assert workbook.sheetnames == ["提取结果", "需人工核对"]
    assert workbook["提取结果"].max_row == 11
    assert workbook["提取结果"]["A2"].value == "0001"
    assert workbook["提取结果"]["A2"].data_type == "s"
    assert workbook["提取结果"]["B2"].value == 113
    assert workbook["提取结果"]["B2"].data_type == "n"
    workbook["提取结果"]["C2"] = "用户补充"
    workbook.save(first)
    before = first.read_bytes()
    second = module.export_results(records, tmp_path, "十张票", ["发票号码", "价税合计", "摘要"])
    assert first != second and first.read_bytes() == before
    assert openpyxl.load_workbook(second)["提取结果"].max_row == 11


def test_核对原因缺失原因与公式样式文本保留(tmp_path):
    module = output()
    records = [record(status="review")]
    records[0]["fields"].update({"摘要": "=SUM(1,2)", "备注": "+123", "扩展": "@名称", "负数字符串": "-ABC"})
    result = module.export_results(records, tmp_path, "核对", ["发票号码", "价税合计", "摘要", "备注", "扩展", "负数字符串"])
    workbook = openpyxl.load_workbook(result, data_only=False)
    assert workbook["提取结果"].max_row == 1
    review = workbook["需人工核对"]
    headers = {cell.value: cell.column for cell in review[1]}
    assert review.cell(2, headers["核对原因"]).value == "金额关系不符"
    assert "购买方" in review.cell(2, headers["取数说明"]).value
    assert "原文未记载" in review.cell(2, headers["取数说明"]).value
    for column, expected in [(3, "=SUM(1,2)"), (4, "+123"), (5, "@名称"), (6, "-ABC")]:
        assert review.cell(2, column).value == expected
        assert review.cell(2, column).data_type == "s"


def create_template(path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "业务台账"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "模板标题"
    sheet["A3"] = "发票号码"
    sheet["C3"] = "价税合计"
    sheet["A4"] = "已有记录"
    sheet["C4"] = 20
    sheet["C5"].number_format = "#,##0.00"
    sheet["A5"].font = Font(name="宋体", bold=True)
    sheet.column_dimensions["A"].width = 24
    sheet.print_title_rows = "1:3"
    instructions = workbook.create_sheet("说明")
    instructions["A1"] = "原有公式"
    instructions["B1"] = "=SUM('业务台账'!C4:C5)"
    workbook.create_sheet("需人工核对")["A1"] = "原有核对内容"
    workbook.save(path)


def test_模板按物理列写入且保留原数据公式与同名表(tmp_path):
    module = output()
    template = tmp_path / "原模板.xlsx"
    create_template(template)
    before = sha256(template.read_bytes()).hexdigest()
    specification = module.inspect_template(template, sheet="业务台账", header_row=3, start_row=5)
    assert specification["fields"] == ["发票号码", "价税合计"]
    assert specification["columns"] == {"发票号码": 1, "价税合计": 3}
    result = module.export_results([record(), record("0002", "review")], tmp_path / "结果", "按模板", specification["fields"], specification)
    assert sha256(template.read_bytes()).hexdigest() == before
    workbook = openpyxl.load_workbook(result, data_only=False)
    sheet = workbook["业务台账"]
    assert sheet["A4"].value == "已有记录" and sheet["C4"].value == 20
    assert sheet["A5"].value == "0001" and sheet["B5"].value is None and sheet["C5"].value == 113
    assert sheet["C5"].number_format == "#,##0.00"
    assert sheet["A5"].font.bold is True
    assert sheet.column_dimensions["A"].width == 24
    assert str(sheet.print_title_rows) == "$1:$3"
    assert "A1:C1" in str(sheet.merged_cells)
    assert workbook["说明"]["B1"].value == "=SUM('业务台账'!C4:C5)"
    assert workbook["需人工核对"]["A1"].value == "原有核对内容"
    assert workbook["需人工核对2"]["A2"].value == "0002"


@pytest.mark.parametrize("occupied", ["已有内容", "=1+2"])
def test_模板指定写入区域已有内容不覆盖(tmp_path, occupied):
    module = output()
    template = tmp_path / "占用模板.xlsx"
    create_template(template)
    workbook = openpyxl.load_workbook(template)
    workbook["业务台账"]["A5"] = occupied
    workbook.save(template)
    before = template.read_bytes()
    with pytest.raises(ValueError, match="已有|公式|写入|空白"):
        spec = module.inspect_template(template, "业务台账", 3, 5)
        module.export_results([record()], tmp_path / "结果", "不覆盖", spec["fields"], spec)
    assert template.read_bytes() == before


@pytest.mark.parametrize("invalid", ["重复表头", "数据合并", "宏文件"])
def test_不支持的模板结构明确拒绝(tmp_path, invalid):
    module = output()
    template = tmp_path / ("复杂模板.xlsm" if invalid == "宏文件" else "复杂模板.xlsx")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["号码", "金额" if invalid != "重复表头" else "号码"])
    if invalid == "数据合并":
        sheet.merge_cells("A2:B2")
    workbook.save(template)
    with pytest.raises(ValueError, match="模板|合并|重复|xlsx"):
        module.inspect_template(template)


def test_模板识别后改变不能继续套用旧定位(tmp_path):
    module = output()
    template = tmp_path / "原模板.xlsx"
    create_template(template)
    spec = module.inspect_template(template, "业务台账", 3, 5)
    workbook = openpyxl.load_workbook(template)
    workbook["业务台账"]["A3"] = "不同字段"
    workbook.save(template)
    with pytest.raises(ValueError, match="改变|变化|重新"):
        module.export_results([record()], tmp_path / "结果", "旧定位", spec["fields"], spec)


def test_超单元格容量明确报错不截断(tmp_path):
    module = output()
    records = [record()]
    records[0]["fields"]["摘要"] = "长" * 32768
    with pytest.raises(ValueError, match="容量|32767|过长"):
        module.export_results(records, tmp_path, "超容量", ["摘要"])

def test_片段进度与来源结果共同保存且重开可续作(tmp_path):
    Store = storage()
    store = Store(tmp_path)
    checkpoint = {"chunks": {"0": {"records": [record()], "complete": True, "unprocessed": []}}, "total": 3}
    store.replace_source("source-a", "多页.pdf", [record()], complete=False, checkpoint=checkpoint)
    store.close()
    store = Store(tmp_path)
    assert store.get_checkpoint("source-a")["total"] == 3
    assert list(store.get_checkpoint("source-a")["chunks"]) == ["0"]
    assert store.get_checkpoint("不存在") is None
    assert store.source_done("source-a") is False
    failure = store.save_attempt("source-a", "服务返回错误原文", error="连接失败")
    assert "连接失败" in Path(failure).read_text(encoding="utf-8-sig")
    store.replace_source("source-a", "多页.pdf", [record(), record("0002")], checkpoint={"chunks": {}, "total": 0})
    assert store.source_done("source-a") is True
    assert store.get_checkpoint("source-a")["total"] == 0
    assert len(store.records()) == 2
    store.close()

def test_模板各行已有不同样式不能被首行样式覆盖(tmp_path):
    module = output()
    template = tmp_path / "不同样式.xlsx"
    create_template(template)
    workbook = openpyxl.load_workbook(template)
    workbook["业务台账"]["A6"].font = Font(name="宋体", italic=True, color="FF0000")
    workbook.save(template)
    spec = module.inspect_template(template, "业务台账", 3, 5)
    result = module.export_results([record(), record("0002")], tmp_path / "结果", "保留样式", spec["fields"], spec)
    worksheet = openpyxl.load_workbook(result)["业务台账"]
    assert worksheet["A5"].font.bold is True
    assert worksheet["A6"].font.italic is True
    assert worksheet["A6"].font.color.rgb == "00FF0000"


def test_复杂绘图模板和超行容量不能静默输出(tmp_path):
    from openpyxl.chart import BarChart, Reference
    module = output()
    template = tmp_path / "图表模板.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(["编号", "金额"])
    worksheet.append(["001", 12])
    chart = BarChart()
    chart.add_data(Reference(worksheet, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    worksheet.add_chart(chart, "E1")
    workbook.save(template)
    before = template.read_bytes()
    with pytest.raises(ValueError, match="绘图|对象|模板"):
        module.inspect_template(template)
    assert template.read_bytes() == before
    plain = tmp_path / "行容量.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["发票号码"])
    workbook.save(plain)
    specification = module.inspect_template(plain, start_row=1048576)
    with pytest.raises(ValueError, match="行数|容量"):
        module.export_results([record(), record("0002")], tmp_path / "结果", "超行", specification["fields"], specification)

def test_超过Excel有效数字的金额不能悄悄舍入(tmp_path):
    module = output()
    with pytest.raises(ValueError, match="精度|有效数字"):
        module.export_results([record(amount=Decimal("123456789012345.67"))], tmp_path, "高精度金额", ["发票号码", "价税合计"])