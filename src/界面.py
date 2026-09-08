"""面向日常办公的Windows桌面窗口。"""
from pathlib import Path
from datetime import datetime
import os
import queue
import threading
import uuid
import tkinter as tk
from tkinter import filedialog, messagebox
import customtkinter as ctk
from openpyxl import load_workbook
from .本机设置 import load_settings, save_settings
from .字段配置 import FIELDS
from .应用流程 import Pipeline
from .模型调用 import Client
from .任务存储 import Store
from .模板输出 import inspect_template


class App(ctk.CTk):
    def __init__(self,root_dir=None):
        super().__init__()
        self.root_dir=Path(root_dir or Path(__file__).resolve().parents[1])
        self.title('文档智能填表系统')
        self.geometry('1080x880');self.minsize(920,700)
        ctk.set_appearance_mode('light')
        self.files=[];self.template_path=None;self.template=None
        self.last_result=None;self.batch_dir=None;self.busy=False;self.closing=False
        self.cancel_event=threading.Event();self.events=queue.Queue()
        self.type_var=tk.StringVar(value='发票')
        self.output_var=tk.StringVar(value=str(self.root_dir/'结果'))
        self.header_var=tk.StringVar(value='1');self.start_var=tk.StringVar(value='')
        self.remember_var=tk.BooleanVar(value=False)
        self._build()
        try:
            saved=load_settings(self.root_dir)
        except ValueError as error:
            saved={};self.log(str(error))
        for entry,key in [(self.url_entry,'base_url'),(self.key_entry,'api_key'),(self.model_entry,'model')]:
            entry.insert(0,saved.get(key,''))
        self.remember_var.set(bool(saved.get('remember_key',False)))
        self.protocol('WM_DELETE_WINDOW',self.close_window)
        self.after(120,self.poll_events)

    def _build(self):
        outer=ctk.CTkScrollableFrame(self,fg_color='#F4F7FB')
        outer.pack(fill='both',expand=True,padx=14,pady=12)
        outer.grid_columnconfigure(0,weight=1)
        ctk.CTkLabel(outer,text='文档智能填表',font=('Microsoft YaHei',26,'bold'),text_color='#174675').grid(row=0,column=0,sticky='w',padx=16,pady=(6,2))
        ctk.CTkLabel(outer,text='选材料 → 提取和核对 → 保存为Excel。原始文件保持不变。',font=('Microsoft YaHei',13),text_color='#566477').grid(row=1,column=0,sticky='w',padx=16,pady=(0,10))
        source=ctk.CTkFrame(outer,fg_color='white');source.grid(row=2,column=0,sticky='ew',pady=5);source.grid_columnconfigure(1,weight=1)
        ctk.CTkLabel(source,text='1  选择材料',font=('Microsoft YaHei',16,'bold')).grid(row=0,column=0,sticky='w',padx=14,pady=10)
        ctk.CTkOptionMenu(source,values=list(FIELDS),variable=self.type_var,width=150).grid(row=0,column=1,sticky='w')
        buttons=ctk.CTkFrame(source,fg_color='transparent');buttons.grid(row=0,column=2,padx=12)
        ctk.CTkButton(buttons,text='选文件',width=100,command=self.select_files).pack(side='left',padx=4)
        ctk.CTkButton(buttons,text='选文件夹',width=100,command=self.select_folder).pack(side='left',padx=4)
        ctk.CTkButton(buttons,text='清空列表',width=90,fg_color='#738297',command=self.clear_files).pack(side='left',padx=4)
        self.file_box=ctk.CTkTextbox(source,height=85,font=('Microsoft YaHei',12));self.file_box.grid(row=1,column=0,columnspan=3,sticky='ew',padx=14,pady=(0,10));self.file_box.configure(state='disabled')
        ctk.CTkLabel(source,text='结果保存位置').grid(row=2,column=0,padx=14,pady=(0,12),sticky='w')
        ctk.CTkEntry(source,textvariable=self.output_var,state='readonly').grid(row=2,column=1,sticky='ew',pady=(0,12))
        ctk.CTkButton(source,text='选择保存位置',width=130,command=self.select_output).grid(row=2,column=2,padx=14,pady=(0,12),sticky='e')
        connect=ctk.CTkFrame(outer,fg_color='white');connect.grid(row=3,column=0,sticky='ew',pady=5);connect.grid_columnconfigure(1,weight=1);connect.grid_columnconfigure(3,weight=1)
        ctk.CTkLabel(connect,text='2  连接模型',font=('Microsoft YaHei',16,'bold')).grid(row=0,column=0,columnspan=4,sticky='w',padx=14,pady=10)
        ctk.CTkLabel(connect,text='服务地址').grid(row=1,column=0,padx=14,pady=5,sticky='w')
        self.url_entry=ctk.CTkEntry(connect,placeholder_text='填写服务提供方给你的连接地址');self.url_entry.grid(row=1,column=1,columnspan=3,sticky='ew',padx=(0,14),pady=5)
        ctk.CTkLabel(connect,text='访问密钥').grid(row=2,column=0,padx=14,pady=5,sticky='w')
        self.key_entry=ctk.CTkEntry(connect,show='●');self.key_entry.grid(row=2,column=1,sticky='ew',pady=5)
        ctk.CTkLabel(connect,text='模型名称').grid(row=2,column=2,padx=14)
        self.model_entry=ctk.CTkEntry(connect,placeholder_text='服务提供方给出的模型名称');self.model_entry.grid(row=2,column=3,sticky='ew',padx=(0,14),pady=5)
        ctk.CTkCheckBox(connect,text='在这台电脑记住密钥',variable=self.remember_var).grid(row=3,column=1,sticky='w',pady=(7,12))
        self.test_button=ctk.CTkButton(connect,text='检查连接',width=130,command=self.test_connection);self.test_button.grid(row=3,column=3,sticky='e',padx=14,pady=(7,12))
        template=ctk.CTkFrame(outer,fg_color='white');template.grid(row=4,column=0,sticky='ew',pady=5);template.grid_columnconfigure(1,weight=1)
        ctk.CTkLabel(template,text='3  表格要求（可不选）',font=('Microsoft YaHei',16,'bold')).grid(row=0,column=0,columnspan=2,sticky='w',padx=14,pady=10)
        ctk.CTkButton(template,text='选择Excel模板',width=135,command=self.select_template).grid(row=0,column=2,padx=8)
        ctk.CTkButton(template,text='使用默认字段',width=135,fg_color='#738297',command=self.clear_template).grid(row=0,column=3,padx=(0,14))
        self.template_label=ctk.CTkLabel(template,text='未选择模板，将使用所选类型的默认字段。',anchor='w',wraplength=900,text_color='#566477');self.template_label.grid(row=1,column=0,columnspan=4,sticky='ew',padx=14)
        self.sheet_menu=ctk.CTkOptionMenu(template,values=['未选模板'],width=150);self.sheet_menu.grid(row=2,column=0,padx=14,pady=8)
        rowframe=ctk.CTkFrame(template,fg_color='transparent');rowframe.grid(row=2,column=1,columnspan=2,sticky='w')
        ctk.CTkLabel(rowframe,text='表头在第').pack(side='left');ctk.CTkEntry(rowframe,textvariable=self.header_var,width=55).pack(side='left',padx=5);ctk.CTkLabel(rowframe,text='行；开始填入第').pack(side='left');ctk.CTkEntry(rowframe,textvariable=self.start_var,width=65,placeholder_text='自动').pack(side='left',padx=5);ctk.CTkLabel(rowframe,text='行').pack(side='left')
        ctk.CTkButton(template,text='预览模板字段',width=135,command=self.preview_template).grid(row=2,column=3,padx=(0,14))
        ctk.CTkLabel(template,text='补充取数说明',anchor='w').grid(row=3,column=0,sticky='w',padx=14)
        self.instructions_entry=ctk.CTkEntry(template,placeholder_text='例如：金额取合同含税总金额；没有记载就留空');self.instructions_entry.grid(row=3,column=1,columnspan=3,sticky='ew',padx=(0,14),pady=(0,12))
        action=ctk.CTkFrame(outer,fg_color='transparent');action.grid(row=5,column=0,sticky='ew',pady=10)
        self.start_button=ctk.CTkButton(action,text='开始处理',height=38,width=160,command=self.start_processing);self.start_button.pack(side='left',padx=(0,8))
        self.stop_button=ctk.CTkButton(action,text='停止',height=38,width=85,state='disabled',fg_color='#A45045',command=self.stop_processing);self.stop_button.pack(side='left',padx=4)
        self.resume_button=ctk.CTkButton(action,text='恢复已有批次',height=38,width=145,command=self.resume_batch);self.resume_button.pack(side='left',padx=4)
        self.retry_button=ctk.CTkButton(action,text='重试需核对文件',height=38,width=150,command=self.retry_batch);self.retry_button.pack(side='left',padx=4)
        self.open_button=ctk.CTkButton(action,text='打开结果',height=38,width=105,command=self.open_output);self.open_button.pack(side='left',padx=4)
        ctk.CTkButton(action,text='查看核对项',height=38,width=110,command=self.show_review).pack(side='left',padx=4)
        self.progress=ctk.CTkProgressBar(outer);self.progress.grid(row=6,column=0,sticky='ew',padx=2);self.progress.set(0)
        self.status_label=ctk.CTkLabel(outer,text='准备就绪。先选择一种文档类型和材料。',anchor='w',font=('Microsoft YaHei',13));self.status_label.grid(row=7,column=0,sticky='ew',pady=4)
        self.log_box=ctk.CTkTextbox(outer,height=140,font=('Microsoft YaHei',12));self.log_box.grid(row=8,column=0,sticky='ew');self.log_box.configure(state='disabled')
        ctk.CTkLabel(outer,text='核对项会单独列出。请核对重要金额、日期和事项后再用于正式工作。',text_color='#566477',anchor='w').grid(row=9,column=0,sticky='w',pady=8)

    def log(self,text):
        self.log_box.configure(state='normal');self.log_box.insert('end',text+'\n');self.log_box.see('end');self.log_box.configure(state='disabled')

    def display_files(self):
        self.file_box.configure(state='normal');self.file_box.delete('1.0','end');self.file_box.insert('end','\n'.join(str(p) for p in self.files));self.file_box.configure(state='disabled')

    def select_files(self):
        if self.busy:return
        paths=filedialog.askopenfilenames(title='选择需要填表的材料',filetypes=[('文档、图片和PDF','*.doc *.docx *.pdf *.png *.jpg *.jpeg *.tif *.tiff')])
        self.files=list(dict.fromkeys(self.files+[Path(p) for p in paths]));self.display_files()

    def select_folder(self):
        if self.busy:return
        p=filedialog.askdirectory(title='选择材料文件夹（包括其中子文件夹）')
        if p:self.files.append(Path(p));self.display_files()

    def clear_files(self):
        if not self.busy:self.files=[];self.display_files()

    def select_output(self):
        if self.busy:return
        p=filedialog.askdirectory(title='选择结果保存位置')
        if p:self.output_var.set(p)

    def select_template(self):
        if self.busy:return
        p=filedialog.askopenfilename(title='选择普通Excel模板',filetypes=[('Excel工作簿','*.xlsx')])
        if not p:return
        try:
            w=load_workbook(p,read_only=True);names=w.sheetnames;w.close()
            self.template_path=Path(p);self.template=None
            self.sheet_menu.configure(values=names);self.sheet_menu.set(names[0])
            self.template_label.configure(text='已选择：'+Path(p).name+'；请确认工作表及表头行，再预览字段。')
        except Exception:
            messagebox.showwarning('模板无法读取','请确认文件是可打开的普通Excel工作簿。')

    def clear_template(self):
        if self.busy:return
        self.template_path=None;self.template=None
        self.template_label.configure(text='未选择模板，将使用所选类型的默认字段。')
        self.sheet_menu.configure(values=['未选模板']);self.sheet_menu.set('未选模板')

    def preview_template(self):
        if not self.template_path:return None
        try:
            value=inspect_template(self.template_path,sheet=self.sheet_menu.get(),header_row=int(self.header_var.get()),start_row=int(self.start_var.get()) if self.start_var.get().strip() else None)
            self.template=value
            self.template_label.configure(text='填写字段：'+'、'.join(value['fields'])+f'。从第{value["start_row"]}行写入，来源信息另附在右侧。')
            return value
        except Exception as error:
            self.template=None;messagebox.showwarning('请检查模板',str(error));return None

    def read_form(self):
        return {'base_url':self.url_entry.get().strip().rstrip('/'),'api_key':self.key_entry.get().strip(),'model':self.model_entry.get().strip(),'remember_key':bool(self.remember_var.get())}

    def set_busy(self,value):
        self.busy=value
        for button in [self.start_button,self.test_button,self.resume_button,self.retry_button]:button.configure(state='disabled' if value else 'normal')
        self.stop_button.configure(state='normal' if value else 'disabled')

    def _settings_valid(self,settings):
        if not all(settings[k] for k in ['base_url','api_key','model']):
            messagebox.showwarning('先填写连接信息','请填写服务地址、访问密钥和模型名称。');return False
        return True

    def start_processing(self):
        if self.busy:return
        if not self.files:
            messagebox.showwarning('先选择材料','请先选择文件或文件夹。');return
        settings=self.read_form()
        if not self._settings_valid(settings):return
        if self.template_path and not self.preview_template():return
        self.batch_dir=self.root_dir/'.local'/'批次'/(datetime.now().strftime('%Y%m%d_%H%M%S')+'_'+uuid.uuid4().hex[:6])
        self._launch(settings,list(self.files))

    def _launch(self,settings,paths,resume=False,retry_only=False):
        try:save_settings(self.root_dir,settings)
        except Exception as error:messagebox.showwarning('设置未保存',str(error));return
        self.cancel_event=threading.Event();self.set_busy(True);self.progress.set(0)
        pipeline=Pipeline(settings,self.type_var.get(),self.output_var.get(),self.batch_dir,template=self.template,instructions=self.instructions_entry.get().strip(),on_event=self.events.put)
        def run():
            try:self.events.put({'result':pipeline.run(paths,resume=resume,retry_only=retry_only,cancel_event=self.cancel_event)})
            except Exception as error:self.events.put({'error':str(error)})
        threading.Thread(target=run,daemon=True).start()

    def resume_batch(self):
        if self.busy:return
        folder=filedialog.askdirectory(title='选择需要恢复的批次文件夹',initialdir=str(self.root_dir/'.local'/'批次'))
        if not folder:return
        self._resume_path(Path(folder),False)

    def _resume_path(self,folder,retry_only):
        try:
            store=Store(folder);config=store.get_config();sources=store.sources();store.close()
            if not config:raise ValueError('这里没有可恢复的批次，请选择本程序创建的批次文件夹。')
            settings=self.read_form()
            if not self._settings_valid(settings):return
            self.batch_dir=folder;self.type_var.set(config['kind']);self.template=config.get('template')
            self.template_path=Path(self.template['path']) if self.template else None
            self.output_var.set(config['output_dir'])
            self.instructions_entry.delete(0,'end');self.instructions_entry.insert(0,config.get('instructions',''))
            self.files=[Path(s['path']) for s in sources];self.display_files()
            self._launch(settings,self.files,resume=True,retry_only=retry_only)
        except Exception as error:messagebox.showwarning('无法恢复',str(error))

    def retry_batch(self):
        if self.busy:return
        if self.batch_dir:self._resume_path(self.batch_dir,True)
        else:messagebox.showwarning('先选择批次','请先恢复一个已有批次，再重试其中需核对的文件。')

    def stop_processing(self):
        self.cancel_event.set();self.status_label.configure(text='正在停止后续处理，已保存结果会保留。')

    def test_connection(self):
        if self.busy:return
        settings=self.read_form()
        if not self._settings_valid(settings):return
        self.cancel_event=threading.Event();self.set_busy(True)
        def run():
            try:
                Client(settings).extract([{'kind':'text','position':'连接示例','text':'收据，2026年9月8日，示例公司收到人民币1元。'}],'发票',cancel_event=self.cancel_event)
                self.events.put({'connection':'连接成功，模型已返回可读取的结果。'})
            except Exception as error:self.events.put({'error':str(error)})
        threading.Thread(target=run,daemon=True).start()

    def poll_events(self):
        try:
            while True:
                event=self.events.get_nowait()
                if 'message' in event:
                    self.log(event['message']);self.status_label.configure(text=event.get('phase',event['message']))
                    if event.get('total'):self.progress.set(event.get('done',0)/event['total'])
                if 'result' in event:
                    self.last_result=event['result'];self.set_busy(False)
                    result=self.last_result
                    suffix='已停止，可恢复。' if result['cancelled'] else ('连接需检查，可恢复。' if result['paused'] else '处理结束。')
                    self.status_label.configure(text=f'{suffix} 正常 {result["normal"]} 行，需核对 {result["review"]} 行。'+('结果已导出。' if result['output'] else 'Excel尚未导出。'))
                    self.log(self.status_label.cget('text'));self.progress.set(1 if result['complete'] else self.progress.get())
                if 'connection' in event:
                    self.set_busy(False);self.log(event['connection']);self.status_label.configure(text=event['connection'])
                if 'error' in event:
                    self.set_busy(False);self.status_label.configure(text='尚未完成，请检查下方说明。');self.log(event['error'])
        except queue.Empty:pass
        if self.closing and not self.busy:
            self.destroy();return
        self.after(120,self.poll_events)

    def open_output(self):
        p=self.last_result.get('output') if self.last_result else None
        if p and Path(p).exists():os.startfile(p)
        else:messagebox.showinfo('尚无结果','处理并导出完成后，可以在这里打开结果Excel。')

    def show_review(self):
        rows=[r for r in (self.last_result or {}).get('rows',[]) if r.get('status')!='normal']
        if not rows:messagebox.showinfo('核对清单','当前没有可显示的核对项，请先处理或恢复批次。');return
        window=ctk.CTkToplevel(self);window.title('需人工核对');window.geometry('850x550')
        box=ctk.CTkTextbox(window,font=('Microsoft YaHei',13));box.pack(fill='both',expand=True,padx=15,pady=15)
        for number,row in enumerate(rows,1):
            box.insert('end',f'{number}. {row.get("source_file", "")}\n位置：'+ '；'.join(row.get('positions',[]))+'\n原因：'+'；'.join(row.get('reasons',[]))+'\n已识别：'+str(row.get('fields',{}))+'\n\n')
        box.configure(state='disabled');window.after(100,window.lift)

    def close_window(self):
        if self.busy:
            self.closing=True;self.stop_processing()
        else:self.destroy()