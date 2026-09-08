"""使用合成原件检查读取、来源定位和切段，不读取客户资料。"""

import base64
import hashlib
import importlib
import io
import uuid
from pathlib import Path

import fitz
import pytest
from docx import Document
from docx.oxml import OxmlElement
from PIL import Image


@pytest.fixture
def sample_dir():
    # 不自动清理：用户要求测试资料也不得永久删除。
    folder = Path(__file__).resolve().parents[1] / "work" / "reader_tests" / uuid.uuid4().hex
    folder.mkdir(parents=True)
    return folder


def read_document(path, work_dir):
    return importlib.import_module("src.文档读取").read_document(path, work_dir)


def chunk_blocks(blocks, **kwargs):
    return importlib.import_module("src.文档读取").chunk_blocks(blocks, **kwargs)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def decoded_image(block):
    assert block["image"].startswith("data:image/png;base64,")
    return Image.open(io.BytesIO(base64.b64decode(block["image"].split(",", 1)[1])))


def test_word_keeps_paragraph_table_image_order_and_source_hash(sample_dir):
    picture = sample_dir / "原图.png"
    Image.new("RGB", (64, 32), "red").save(picture)
    source = sample_dir / "顺序.docx"
    document = Document()
    document.add_paragraph("段落一")
    document.add_table(rows=1, cols=2).rows[0].cells[0].text = "表格左格"
    document.tables[0].cell(0, 1).text = "表格右格"
    paragraph = document.add_paragraph("图前")
    paragraph.add_run().add_picture(str(picture))
    paragraph.add_run("图后")
    document.add_paragraph("末段")
    document.save(source)
    before = digest(source)

    blocks = read_document(source, sample_dir / "工作")

    ordered = [b["text"] if b["kind"] == "text" else "图片" for b in blocks]
    assert ordered == ["段落一", "表格左格", "表格右格", "图前", "图片", "图后", "末段"]
    assert "表格1第1行第2列" in blocks[2]["position"]
    assert decoded_image(blocks[4]).size == (64, 32)
    assert digest(source) == before


def test_word_headers_footers_nested_tables_and_shared_image(sample_dir):
    picture = sample_dir / "标记.png"
    Image.new("RGB", (80, 40), "blue").save(picture)
    document = Document()
    document.add_paragraph("正文")
    document.add_picture(str(picture))
    document.add_picture(str(picture))
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "嵌套表格"
    document.sections[0].header.paragraphs[0].text = "页眉单位"
    document.sections[0].footer.paragraphs[0].text = "页脚编号"
    source = sample_dir / "页眉页脚.docx"
    document.save(source)
    before = digest(source)

    blocks = read_document(source, sample_dir / "工作")

    assert sum(b["kind"] == "image" for b in blocks) == 1
    assert any("重复图片" in b["text"] and b["position"] for b in blocks)
    assert any(b["text"] == "页眉单位" and "页眉" in b["position"] for b in blocks)
    assert any(b["text"] == "页脚编号" and "页脚" in b["position"] for b in blocks)
    assert any(b["text"] == "嵌套表格" and "表格" in b["position"] for b in blocks)
    assert digest(source) == before


def test_word_unsupported_embedded_region_is_reported(sample_dir):
    document = Document()
    paragraph = document.add_paragraph("可读正文")
    paragraph.add_run()._r.append(OxmlElement("w:object"))
    source = sample_dir / "嵌入对象.docx"
    document.save(source)
    blocks = read_document(source, sample_dir / "工作")
    assert any(b["kind"] == "warning" and "段落1" in b["position"] and "嵌入对象" in b["text"] for b in blocks)
    assert any(b["text"] == "可读正文" for b in blocks)


def test_pdf_renders_all_pages_including_same_page_pictures(sample_dir):
    source = sample_dir / "两页.pdf"
    pdf = fitz.open()
    page = pdf.new_page(width=200, height=120)
    page.insert_text((10, 18), "Text invoice")
    page.draw_rect(fitz.Rect(10, 25, 90, 110), color=(1, 0, 0), fill=(1, 0, 0))
    page.draw_rect(fitz.Rect(110, 25, 190, 110), color=(0, 0, 1), fill=(0, 0, 1))
    page = pdf.new_page(width=200, height=120)
    raster = io.BytesIO()
    Image.new("RGB", (200, 120), "green").save(raster, format="PNG")
    page.insert_image(page.rect, stream=raster.getvalue())
    pdf.save(source)
    pdf.close()
    before = digest(source)

    blocks = read_document(source, sample_dir / "工作")

    pictures = [b for b in blocks if b["kind"] == "image"]
    assert len(pictures) == 2
    assert "第1页" in pictures[0]["position"] and "第2页" in pictures[1]["position"]
    first = decoded_image(pictures[0]).convert("RGB")
    red = first.getpixel((first.width // 4, first.height // 2))
    blue = first.getpixel((first.width * 3 // 4, first.height // 2))
    assert red[0] > 230 and red[2] < 20
    assert blue[2] > 230 and blue[0] < 20
    assert digest(source) == before


def test_multiframe_tiff_keeps_each_page_and_small_image_resolution(sample_dir):
    source = sample_dir / "三页.tiff"
    pages = [Image.new("RGB", (90, 60), color) for color in ("red", "green", "blue")]
    pages[0].save(source, save_all=True, append_images=pages[1:])
    before = digest(source)
    blocks = read_document(source, sample_dir / "工作")
    assert len(blocks) == 3
    assert [decoded_image(b).getpixel((45, 30)) for b in blocks] == [(255, 0, 0), (0, 128, 0), (0, 0, 255)]
    assert all(decoded_image(b).size == (90, 60) for b in blocks)
    assert "第3页" in blocks[2]["position"]
    assert digest(source) == before


@pytest.mark.parametrize("suffix", [".png", ".jpg", ".jpeg"])
def test_static_images_are_read_without_source_changes(sample_dir, suffix):
    source = sample_dir / ("原件" + suffix)
    Image.new("RGB", (100, 70), "white").save(source)
    before = digest(source)
    blocks = read_document(source, sample_dir / "工作")
    assert len(blocks) == 1
    assert decoded_image(blocks[0]).size == (100, 70)
    assert digest(source) == before


def test_large_image_is_partitioned_without_losing_edge_pixels(sample_dir):
    source = sample_dir / "宽图.png"
    picture = Image.new("RGB", (3000, 100), "white")
    picture.putpixel((0, 0), (255, 0, 0))
    picture.putpixel((2999, 99), (0, 0, 255))
    picture.save(source)
    blocks = read_document(source, sample_dir / "工作")
    images = [decoded_image(b) for b in blocks if b["kind"] == "image"]
    assert len(images) > 1
    assert all(im.width <= 2400 and im.height <= 2400 for im in images)
    assert images[0].getpixel((0, 0)) == (255, 0, 0)
    assert images[-1].getpixel((images[-1].width - 1, images[-1].height - 1)) == (0, 0, 255)
    assert all("区域" in b["position"] for b in blocks)


@pytest.mark.parametrize("suffix", [".docx", ".pdf", ".tiff", ".jpg"])
def test_corrupt_file_has_chinese_reason_and_is_unchanged(sample_dir, suffix):
    source = sample_dir / ("损坏" + suffix)
    source.write_bytes(b"broken input")
    before = digest(source)
    with pytest.raises(ValueError, match="无法读取|文件损坏"):
        read_document(source, sample_dir / "工作")
    assert digest(source) == before


def test_oversized_text_is_fully_covered_with_context_and_bounded_chunks():
    original = "".join(chr(0x4E00 + i) for i in range(700))
    blocks = [{"kind": "text", "position": "段落1", "text": original}]
    chunks = chunk_blocks(blocks, max_chars=120, max_images=2)
    assert len(chunks) > 1
    texts = [b["text"] for c in chunks for b in c]
    assert set("".join(texts)) == set(original)
    assert all(text in original for text in texts)
    assert all(sum(len(b["text"]) for b in c) <= 120 for c in chunks)
    assert all("段落1" in b["position"] for c in chunks for b in c)
    assert any(set(texts[i]) & set(texts[i + 1]) for i in range(len(texts) - 1))
    assert blocks == [{"kind": "text", "position": "段落1", "text": original}]


def test_page_chunking_respects_image_limit_and_keeps_neighbor_page():
    blocks = [{"kind": "image", "position": f"第{i}页", "text": "", "image": f"image-{i}"} for i in range(1, 8)]
    chunks = chunk_blocks(blocks, max_chars=50, max_images=3)
    assert all(sum(b["kind"] == "image" for b in c) <= 3 for c in chunks)
    assert {b["position"] for c in chunks for b in c} == {f"第{i}页" for i in range(1, 8)}
    assert any(chunks[i][-1]["position"] == chunks[i + 1][0]["position"] for i in range(len(chunks) - 1))
    assert len(chunk_blocks(blocks, max_images=1)) == 7


def test_chunking_keeps_warnings_and_rejects_impossible_budget():
    blocks = [{"kind": "warning", "position": "表格2", "text": "无法读取嵌入对象" * 20}]
    chunks = chunk_blocks(blocks, max_chars=40, max_images=1)
    assert chunks and all(b["kind"] == "warning" for c in chunks for b in c)
    assert all(sum(len(b["text"]) for b in c) <= 40 for c in chunks)
    with pytest.raises(ValueError, match="预算"):
        chunk_blocks(blocks, max_chars=0)


def test_doc_conversion_uses_work_copies_and_never_overwrites(sample_dir):
    import pythoncom
    import win32com.client

    source_docx = sample_dir / "旧格式样本.docx"
    source_doc = sample_dir / "旧格式样本.doc"
    document = Document()
    document.add_paragraph("只读转换测试正文")
    document.save(source_docx)
    pythoncom.CoInitialize()
    word = None
    opened = None
    try:
        try:
            word = win32com.client.DispatchEx("Word.Application")
        except Exception as exc:
            pytest.skip(f"本机未能启动 Word，DOC 实机验证未执行：{exc}")
        word.Visible = False
        word.DisplayAlerts = 0
        opened = word.Documents.Open(str(source_docx), ReadOnly=True, AddToRecentFiles=False)
        opened.SaveAs2(str(source_doc), FileFormat=0)
    finally:
        if opened is not None:
            opened.Close(SaveChanges=False)
        if word is not None:
            word.Quit(SaveChanges=False)
        pythoncom.CoUninitialize()
    before = digest(source_doc)
    work_dir = sample_dir / "工作"

    first = read_document(source_doc, work_dir)
    converted_before = {p: digest(p) for p in work_dir.rglob("*.docx")}
    second = read_document(source_doc, work_dir)

    assert any(b["text"] == "只读转换测试正文" for b in first)
    assert any(b["text"] == "只读转换测试正文" for b in second)
    assert len(list(work_dir.rglob("*.docx"))) == 2
    assert all(digest(p) == value for p, value in converted_before.items())
    assert digest(source_doc) == before
    assert not list(sample_dir.glob("~$*"))

def test_word_table_content_controls_do_not_hide_rows_or_cells(sample_dir):
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "被包装的左格"
    table.cell(0, 1).text = "被包装的右格"
    row = table._tbl.tr_lst[0]
    cell = row.tc_lst[0]
    cell_control = OxmlElement("w:sdt")
    cell_content = OxmlElement("w:sdtContent")
    row.replace(cell, cell_control)
    cell_control.append(cell_content)
    cell_content.append(cell)
    row_control = OxmlElement("w:sdt")
    row_content = OxmlElement("w:sdtContent")
    table._tbl.replace(row, row_control)
    row_control.append(row_content)
    row_content.append(row)
    source = sample_dir / "内容控件表格.docx"
    document.save(source)
    blocks = read_document(source, sample_dir / "工作")
    assert [b["text"] for b in blocks if b["kind"] == "text"] == ["被包装的左格", "被包装的右格"]
    assert "第1行第2列" in blocks[1]["position"]


def test_word_alternate_representation_is_not_read_twice(sample_dir):
    from lxml import etree

    document = Document()
    paragraph = document.add_paragraph()
    alternate = etree.fromstring(b'<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><mc:Choice Requires="w"><w:r><w:t>primary</w:t></w:r></mc:Choice><mc:Fallback><w:r><w:t>fallback</w:t></w:r></mc:Fallback></mc:AlternateContent>')
    paragraph._p.append(alternate)
    source = sample_dir / "兼容显示.docx"
    document.save(source)
    blocks = read_document(source, sample_dir / "工作")
    assert [b["text"] for b in blocks if b["kind"] == "text"] == ["primary"]