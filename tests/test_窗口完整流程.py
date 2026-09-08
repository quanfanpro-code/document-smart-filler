"""通过真实Windows控件、HTTP接口、读取与导出验证开始、停止和恢复。"""
import hashlib
import json
import threading
import time
from contextlib import contextmanager
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


def _invoice(number):
    return {"document_id": number, "fields": {
        "票种": "增值税普通发票", "发票号码": number, "开票日期": "2026-09-08",
        "销售方": "示例销售方", "销售方纳税人识别号": "001234567890123456",
        "购买方": "示例购买方", "购买方纳税人识别号": "009876543210987654",
        "不含税金额": "100.00", "税额": "13.00", "价税合计": "113.00",
        "币种": "人民币", "金额单位": "元", "商品/服务明细": "办公用品"},
        "positions": ["段落1"], "issues": [], "missing": {}, "precision": "0.01"}


@contextmanager
def _controlled_service():
    state = {"requests": [], "active": 0}
    second_entered, release_second = threading.Event(), threading.Event()
    lock = threading.Lock()

    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *args):
            pass

        def do_POST(self):
            payload = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            with lock:
                state["requests"].append(payload)
                attempt = len(state["requests"])
                state["active"] += 1
            try:
                text = json.dumps(payload["messages"][-1], ensure_ascii=False)
                number = "GUI002" if "GUI002" in text else "GUI001"
                if attempt == 2:
                    second_entered.set()
                    assert release_second.wait(10), "测试等待停止动作超时"
                content = {"records": [_invoice(number)], "complete": True, "unprocessed": []}
                body = {"id": f"controlled-{attempt}", "object": "chat.completion",
                        "created": 1, "model": payload["model"],
                        "choices": [{"index": 0, "message": {"role": "assistant",
                         "content": json.dumps(content, ensure_ascii=False)}, "finish_reason": "stop"}],
                        "usage": {"prompt_tokens": 30, "completion_tokens": 40, "total_tokens": 70}}
                data = json.dumps(body, ensure_ascii=False).encode("utf-8")
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(data)))
                self.end_headers()
                self.wfile.write(data)
            finally:
                with lock:
                    state["active"] -= 1

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    worker = threading.Thread(target=server.serve_forever, daemon=True)
    worker.start()
    try:
        yield f"http://127.0.0.1:{server.server_port}/v1", state, second_entered, release_second
    finally:
        release_second.set()
        server.shutdown()
        server.server_close()
        worker.join(timeout=2)


def _button(root, text):
    import customtkinter as ctk
    for child in root.winfo_children():
        if isinstance(child, ctk.CTkButton) and child.cget("text") == text:
            return child
        found = _button(child, text)
        if found is not None:
            return found
    return None


def _pump(app, predicate, timeout=12):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        app.update()
        if predicate():
            return
        time.sleep(0.015)
    raise AssertionError("窗口未在限定时间内完成动作：" + app.status_label.cget("text"))


def test_真实窗口按钮完成提取停止恢复和打开结果(tmp_path, monkeypatch):
    from src.界面 import App

    materials = tmp_path / "合成材料"
    materials.mkdir()
    source_paths = []
    for number in ("GUI001", "GUI002"):
        source = materials / (number + ".docx")
        document = Document()
        document.add_paragraph(
            f"增值税普通发票，发票号码{number}，2026年9月8日，示例销售方销售办公用品给示例购买方；"
            "不含税人民币100.00元，税额13.00元，价税合计113.00元。")
        document.save(source)
        source_paths.append(source)
    before = [hashlib.sha256(p.read_bytes()).hexdigest() for p in source_paths]
    output_dir = tmp_path / "结果"
    output_dir.mkdir()
    app_root = tmp_path / "独立应用配置"
    app = App(root_dir=app_root)
    app.geometry("1100x880+50+50")
    try:
        with _controlled_service() as (url, state, second_entered, release_second):
            # 只替换Windows选择框和打开Excel的外部边界，产品窗口、Client、Pipeline均保持真实。
            monkeypatch.setattr("src.界面.filedialog.askopenfilenames", lambda **kw: tuple(map(str, source_paths)))
            monkeypatch.setattr("src.界面.filedialog.askdirectory", lambda **kw: str(output_dir))
            opened = []
            monkeypatch.setattr("src.界面.os.startfile", opened.append)
            app.update()
            assert app.winfo_viewable()
            _button(app, "选文件").invoke()
            _button(app, "选择保存位置").invoke()
            assert app.files == source_paths and app.output_var.get() == str(output_dir)
            for entry, value in ((app.url_entry, url), (app.key_entry, "fake-gui-key"),
                                 (app.model_entry, "窗口受控模型")):
                entry.delete(0, "end")
                entry.insert(0, value)
            app.start_button.invoke()
            assert app.busy and app.start_button.cget("state") == "disabled"
            batch = app.batch_dir
            app.start_button.invoke()
            assert app.batch_dir == batch
            _pump(app, second_entered.is_set)
            assert len(state["requests"]) == 2
            assert state["requests"][0]["model"] == "窗口受控模型"
            assert state["requests"][0]["chat_template_kwargs"]["enable_thinking"] is False
            app.stop_button.invoke()
            _pump(app, lambda: not app.busy and app.last_result is not None)
            stopped = app.last_result
            assert stopped["cancelled"] and not stopped["complete"]
            assert stopped["normal"] == 1 and stopped["review"] == 1
            assert Path(stopped["output"]).is_file()
            release_second.set()
            _pump(app, lambda: state["active"] == 0)
            monkeypatch.setattr("src.界面.filedialog.askdirectory", lambda **kw: str(batch))
            app.resume_button.invoke()
            assert app.busy
            _pump(app, lambda: not app.busy and app.last_result is not stopped)
            finished = app.last_result
            assert finished["complete"] and not finished["cancelled"] and not finished["paused"]
            assert finished["normal"] == 2 and finished["review"] == 0
            assert len(state["requests"]) == 3, "恢复不应重问已保存的第一份文件"
            assert len(list((app_root / ".local" / "批次").iterdir())) == 1
            assert [row["fields"]["发票号码"] for row in finished["rows"]] == ["GUI001", "GUI002"]
            result_path = Path(finished["output"])
            assert result_path.is_file()
            workbook = load_workbook(result_path, data_only=False)
            try:
                sheet = workbook["提取结果"]
                headers = [cell.value for cell in sheet[1]]
                column = headers.index("发票号码") + 1
                assert [sheet.cell(row, column).value for row in (2, 3)] == ["GUI001", "GUI002"]
                assert sheet.max_row == 3
                assert workbook["需人工核对"].max_row == 1
            finally:
                workbook.close()
            app.open_button.invoke()
            assert opened == [str(result_path)]
            saved = json.loads((app_root / ".local" / "设置.json").read_text(encoding="utf-8-sig"))
            assert saved["api_key"] == "", "默认不得持久化用户输入的密钥"
            assert [hashlib.sha256(p.read_bytes()).hexdigest() for p in source_paths] == before
    finally:
        if app.busy:
            app.stop_processing()
            _pump(app, lambda: not app.busy)
        app.destroy()
